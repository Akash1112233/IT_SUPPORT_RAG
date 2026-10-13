
from pathlib import Path
from typing import Iterable

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path) -> list[Document]:
    """Load a supported file and normalize its source metadata."""
    path = Path(path).resolve()
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        documents = PyPDFLoader(str(path)).load()

    elif suffix in {".txt", ".md"}:
        documents = TextLoader(
            str(path),
            encoding="utf-8",
        ).load()

    elif suffix == ".docx":
        doc = DocxDocument(str(path))
        content = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        )
        documents = [
            Document(
                page_content=content,
                metadata={"source": str(path)},
            )
        ]

    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    # Preserve loader metadata such as PDF page numbers, while
    # giving every source a consistent document identity.
    for document in documents:
        document.metadata["source"] = str(path)
        document.metadata["document_name"] = path.name

    return documents


def chunk_documents(
    docs: Iterable[Document],
) -> list[Document]:
    """Split documents and attach consistent chunk metadata."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900,
        chunk_overlap=120,
        add_start_index=True,
    )

    chunks = splitter.split_documents(list(docs))

    for index, chunk in enumerate(chunks):
        source = chunk.metadata.get("source", "unknown")
        chunk.metadata["document_name"] = chunk.metadata.get(
            "document_name",
            Path(str(source)).name,
        )
        chunk.metadata["chunk_index"] = index

    return chunks
